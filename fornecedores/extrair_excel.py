"""Extrai a planilha "Estruturação de Fornecedores" para CSVs normalizados,
prontos para importar nas listas do SharePoint (ver provisionar-listas.ps1).

Uso:  python3 extrair_excel.py "Estruturação de Fornecedores.xlsx"
Gera: dados/*.csv  e imprime um resumo dos índices calculados, para conferir
      com a aba DASHBOARD da planilha.
"""
import csv
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string as col

OUT = Path(__file__).parent / "dados"

SEGMENTOS = [
    "Usinagem",
    "Caldeiraria",
    "Pintura",
    "Tratamento Térmico/Superficial",
    "Prototipagem (Compósito)",
]
# coluna inicial de cada segmento em cada aba
COL_CAPACIDADE = ["B", "F", "J", "N", "R"]
COL_QUALIFICACAO = ["B", "I", "P", "W", "AD"]


def num(v, peso=None):
    """Converte a nota da célula; '-' vira 0, '=Cx' (nota = peso) vira o peso."""
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        if v == "-":
            return 0.0
        if v.startswith("="):
            return peso
        return None
    return float(v)


def status_txt(v):
    """'Status: fornecedor homologado' -> 'Homologado'"""
    if not v or ":" not in v:
        return "Em avaliação"
    return v.split(":", 1)[1].strip().removeprefix("fornecedor ").capitalize()


def write(name, header, rows):
    with open(OUT / name, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(header)
        w.writerows(rows)


def main(path):
    wb = openpyxl.load_workbook(path)
    OUT.mkdir(exist_ok=True)
    avisos = []

    # ---------- Cadastro (contatos) ----------
    ws = wb["CADASTRO DE FORNECEDORES"]
    contatos = {}
    seg_cadastro = {}
    secao = None
    for r in range(1, ws.max_row + 1):
        a, nome = ws.cell(r, 1).value, ws.cell(r, 2).value
        if isinstance(a, str):
            secao = next(s for s in SEGMENTOS if s.upper().startswith(a.strip()))
        if nome and nome != "Empresa":
            contatos[nome.strip()] = [ws.cell(r, c).value or "" for c in (3, 4, 5)]
            seg_cadastro[nome.strip()] = secao

    # ---------- Capacidade técnica ----------
    ws = wb["CAPACIDADE TÉCNICA"]
    catalogo = {}            # (segmento, capacidade) -> peso
    vinculos = {}            # (fornecedor, segmento) -> status
    avaliacoes = []          # fornecedor, segmento, capacidade, peso, nota
    for seg, letra in zip(SEGMENTOS, COL_CAPACIDADE):
        c = col(letra)
        r = 1
        while r <= ws.max_row:
            if str(ws.cell(r, c).value or "").startswith("Fornecedor"):
                forn = ws.cell(r, c + 1).value.strip()
                r += 2  # pula cabeçalho Capacidade/Peso/Nota
                while ws.cell(r, c).value != "TOTAL":
                    cap = ws.cell(r, c).value.strip()
                    # normaliza variações do mesmo item ("Dobra (até 6mm)" x "Dobra")
                    cap = {"Dobra (até 6mm)": "Dobra", "Calandra (até 12mm)": "Calandra",
                           "Eletroerosão (penetração)": "Eletroerosão",
                           "Fudição": "Fundição"}.get(cap, cap)
                    peso = float(ws.cell(r, c + 1).value)
                    catalogo.setdefault((seg, cap), peso)
                    nota = num(ws.cell(r, c + 2).value, peso)
                    if nota is not None:
                        if nota > peso:
                            avisos.append(f"{forn} / {cap}: nota {nota} maior que o peso {peso} — limitada ao peso")
                            nota = peso
                        avaliacoes.append([forn, seg, cap, peso, nota])
                    r += 1
                vinculos[(forn, seg)] = status_txt(ws.cell(r + 1, c).value)
            r += 1

    # fornecedor que só está no Cadastro entra como "em avaliação", sem capacidade
    for nome, seg in seg_cadastro.items():
        if not any(f == nome for f, _ in vinculos):
            vinculos[(nome, seg)] = "Em avaliação"
            avisos.append(f"{nome}: está no Cadastro mas não tem capacidade avaliada")

    # ---------- Critérios de qualificação + histórico ----------
    ws = wb["QUALIFICAÇÃO (REGISTROS)"]
    criterios = []
    historico = []
    for seg, letra in zip(SEGMENTOS, COL_QUALIFICACAO):
        c = col(letra)
        for r in range(9, 15):
            for dim, off in (("Qualidade", 0), ("Prazo", 2), ("Custo", 4)):
                desc, nota = ws.cell(r, c + off).value, ws.cell(r, c + off + 1).value
                if dim != "Qualidade" and seg != SEGMENTOS[0]:
                    continue  # prazo e custo são iguais em todos os segmentos: gravados uma vez só
                if desc is None:
                    avisos.append(f"{seg}: critério de {dim} nota {nota} sem descrição")
                    continue
                criterios.append(["" if dim != "Qualidade" else seg, dim, nota, desc])
        r = 1
        while r <= ws.max_row:
            if ws.cell(r, c).value == "Fornecedor:":
                forn = ws.cell(r, c + 1).value.strip()
                r += 3
                while ws.cell(r, c).value != "Índices Acumulados":
                    d = ws.cell(r, c).value
                    if d is not None:
                        q, p, cu = (ws.cell(r, c + i).value for i in (3, 4, 5))
                        historico.append([forn, seg, d.date().isoformat(), ws.cell(r, c + 1).value,
                                          ws.cell(r, c + 2).value, q, p, cu])
                    r += 1
            r += 1

    # ---------- Grava CSVs ----------
    fornecedores = sorted({f for f, _ in vinculos})
    write("Segmentos.csv", ["Title", "Ordem"], [[s, i + 1] for i, s in enumerate(SEGMENTOS)])
    write("Fornecedores.csv", ["Title", "Telefone", "Endereco", "Responsavel"],
          [[f, *contatos.get(f, ["", "", ""])] for f in fornecedores])
    write("Capacidades.csv", ["Title", "Segmento", "Peso", "Ordem"],
          [[cap, seg, peso, i + 1] for i, ((seg, cap), peso) in enumerate(catalogo.items())])
    write("FornecedorSegmento.csv", ["Fornecedor", "Segmento", "Status"],
          [[f, s, st] for (f, s), st in vinculos.items()])
    write("AvaliacaoCapacidade.csv", ["Fornecedor", "Segmento", "Capacidade", "Nota"],
          [[f, s, cap, nota] for f, s, cap, _, nota in avaliacoes])
    write("CriteriosQualificacao.csv", ["Segmento", "Dimensao", "Nota", "Title"], criterios)
    write("HistoricoServicos.csv",
          ["Fornecedor", "Segmento", "DataOrcamento", "Projeto", "Projetista",
           "NotaQualidade", "NotaPrazo", "NotaCusto"], historico)

    sem_contato = [f for f in fornecedores if f not in contatos]
    avisos.append(f"{len(sem_contato)} fornecedores sem dados de contato no Cadastro: {', '.join(sem_contato)}")

    # ---------- Resumo (mesma conta do DASHBOARD) ----------
    cap_soma, peso_soma = defaultdict(float), defaultdict(float)
    for f, s, cap, peso, nota in avaliacoes:
        cap_soma[(f, s)] += nota
    for (s, cap), peso in catalogo.items():
        peso_soma[s] += peso
    hist = defaultdict(list)
    for f, s, *_, q, p, cu in historico:
        hist[(f, s)].append((q, p, cu))
    print(f"{'Fornecedor':<18}{'Segmento':<32}{'Cap.':>6}{'Qual':>6}{'Prazo':>6}{'Custo':>6}{'Global':>7}{'Serv.':>6}  Status")
    for (f, s), st in vinculos.items():
        h = hist.get((f, s), [])
        med = [sum(x[i] for x in h) / len(h) for i in range(3)] if h else [None] * 3
        glob = sum(med) / 3 if h else None
        fmt = lambda v: f"{v:6.2f}" if v is not None else "     –"
        capn = cap_soma[(f, s)] / peso_soma[s] * 5 if (f, s) in cap_soma else None
        print(f"{f:<18}{s:<32}{fmt(capn)}{fmt(med[0])}{fmt(med[1])}{fmt(med[2])} {fmt(glob)}{len(h):>6}  {st}")
    print("\nAvisos:")
    for a in avisos:
        print(" -", a)


if __name__ == "__main__":
    main(sys.argv[1])
